"""
Excel テレアポリスト → Supabase インポートスクリプト
ヘッダー名で列を自動検出する（シートごとの列順の違いに対応）
"""

import json
import uuid
import urllib.request
import re
import openpyxl
from datetime import datetime

SUPABASE_URL = 'https://antqewvlzfonsakgndxt.supabase.co'
SUPABASE_KEY = 'sb_publishable__37RI5zv54AI023pjvFdbA_AUkoAme2'
EXCEL_PATH   = '/Users/ogisokeisuke/Downloads/【営業リスト_260317~】worktalkテレアポ営業リスト (1).xlsx'

STATUS_MAP = {
    '受けブロ': '受付ブロック', '受付ブロック': '受付ブロック',
    '不在': '担当者不在', '担当者不在': '担当者不在',
    '不通': '不通',
    '担当者接触': '担当者接触',
    '担当者お断り': '断り', '断り': '断り',
    '番号不明/間違い': '不通', '番号不明': '不通',
    '折り返し待ち': '折り返し依頼', '折り返し依頼': '折り返し依頼',
    'アポ取得': 'アポ獲得', 'アポ獲得': 'アポ獲得',
    '問い合わせ': '資料送付', '資料送付': '資料送付',
}

SCALE_MAP = {
    '1〜30名未満': '1〜30名', '30名〜100名未満': '31〜100名',
    '100名〜300名未満': '101〜300名', '100名〜200名未満': '101〜300名',
    '200名〜300名未満': '101〜300名',
    '300名〜500名未満': '301〜500名',
    '500名〜1000名未満': '501〜1000名',
    '1000名〜3000名未満': '1001〜3000名',
    '3000名〜': '3001名〜', '3000名以上': '3001名〜',
}

def normalize_scale(v):
    if not v: return ''
    v = str(v).strip()
    return SCALE_MAP.get(v, v)

def normalize_phone(v):
    if not v: return ''
    return re.sub(r'[\s\-−ー()（）　]', '', str(v))

def to_iso(v):
    if not v: return ''
    if isinstance(v, datetime): return v.isoformat()
    v = str(v).strip()
    for fmt in ['%Y/%m/%d %H:%M:%S', '%Y/%m/%d %H:%M', '%Y/%m/%d', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d']:
        try: return datetime.strptime(v, fmt).isoformat()
        except ValueError: pass
    return v

def find_header_row(ws):
    """企業名 or 会社名を含む行をヘッダー行とする"""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        if any(c in ('企業名', '会社名') for c in row if c):
            return i, list(row)
    return None, None

def build_col_map(headers):
    """ヘッダー名 → 列インデックス のマップを作成"""
    col = {}
    for i, h in enumerate(headers):
        if h is None: continue
        h = str(h).strip()
        if h in ('企業名', '会社名'):   col['company']   = i
        elif h == '業種':              col['industry']  = i
        elif h == '従業員規模':         col['scale']     = i
        elif h in ('番号', '電話番号', '代表電話', 'TEL'): col['phone'] = i
        elif h in ('担当者', '担当者名', '担当者名（先方）'): col['contact'] = i
        elif h in ('架電結果', 'ステータス', '結果'): col['status'] = i
        elif h == 'メモ':              col['memo']      = i
    return col

def extract_call_history(headers, row):
    """架電回数ごとの (TRUE/FALSE, 架電者, 日程) トリプレットを callHistory に変換"""
    history = []
    i = 0
    while i < len(headers):
        h = str(headers[i]).strip() if headers[i] else ''
        # 「1回目」「2回目架電」「11回目」などを検出
        if re.match(r'\d+回目', h):
            result_bool = row[i] if i < len(row) else None
            caller      = str(row[i+1]).strip() if i+1 < len(row) and row[i+1] else ''
            date_val    = row[i+2] if i+2 < len(row) else None

            # TRUE（架電あり）かつ日付 or 架電者がある場合のみ記録
            called = result_bool is True or str(result_bool).upper() == 'TRUE'
            if called and (caller or date_val):
                history.append({
                    'id': str(uuid.uuid4()),
                    'result': '',      # 1回ごとの結果は元データになし
                    'calledBy': caller,
                    'calledAt': to_iso(date_val),
                    'callType': '',
                    'callContent': '',
                    'rejectionReason': '',
                    'note': '',
                })
            i += 3
        else:
            i += 1

    return history

def fetch_existing():
    """既存レコードの企業名・電話番号セットを取得（重複チェック用）"""
    req = urllib.request.Request(
        SUPABASE_URL + '/rest/v1/teleapo_items?select=id,data',
        headers={'apikey': SUPABASE_KEY, 'Authorization': 'Bearer ' + SUPABASE_KEY}
    )
    res = urllib.request.urlopen(req)
    records = json.loads(res.read())
    names, phones = set(), set()
    for r in records:
        d = r['data'] if isinstance(r['data'], dict) else json.loads(r['data'])
        if d.get('companyName'): names.add(d['companyName'].strip().lower())
        if d.get('phone'):       phones.add(normalize_phone(d['phone']))
    return names, phones

def import_sheet(sheet_name, dry_run=False):
    print(f"\n{'='*50}")
    print(f"シート: {sheet_name}")
    print('='*50)

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True, read_only=True)
    if sheet_name not in wb.sheetnames:
        print(f"シート '{sheet_name}' が見つかりません")
        return

    ws = wb[sheet_name]
    header_row_num, headers = find_header_row(ws)
    if not headers:
        print("ヘッダー行が見つかりません")
        return

    col = build_col_map(headers)
    print(f"ヘッダー検出: {col}")

    if 'company' not in col:
        print("企業名列が見つかりません")
        return

    # 既存データ取得
    print("既存データを取得中...")
    existing_names, existing_phones = fetch_existing()
    print(f"既存レコード: {len(existing_names)}件")

    new_items, skipped_dup, skipped_empty = [], 0, 0

    for row in ws.iter_rows(min_row=header_row_num + 1, values_only=True):
        company = str(row[col['company']]).strip() if col.get('company') is not None and row[col['company']] else ''
        if not company or company == 'None':
            skipped_empty += 1
            continue

        phone = normalize_phone(row[col['phone']] if col.get('phone') is not None else '')

        # 重複チェック
        if company.lower() in existing_names or (phone and phone in existing_phones):
            skipped_dup += 1
            continue

        status_raw = str(row[col['status']]).strip() if col.get('status') is not None and row[col['status']] else ''
        memo       = str(row[col['memo']]).strip()   if col.get('memo')   is not None and row[col['memo']]   else ''
        contact    = str(row[col['contact']]).strip() if col.get('contact') is not None and row[col['contact']] else ''
        industry   = str(row[col['industry']]).strip() if col.get('industry') is not None and row[col['industry']] else ''
        scale      = normalize_scale(row[col['scale']] if col.get('scale') is not None else '')

        call_history = extract_call_history(headers, row)

        # アイテムステータス
        mapped_status = STATUS_MAP.get(status_raw, status_raw)
        if mapped_status == 'アポ獲得':
            item_status = 'アポ確定'
        elif mapped_status == '折り返し依頼':
            item_status = '折り返し待ち'
        elif call_history:
            item_status = '架電済'
        else:
            item_status = '未架電'

        # 最後の架電の result にステータスをセット
        if call_history and mapped_status:
            call_history[-1]['result'] = mapped_status

        item = {
            'id': str(uuid.uuid4()),
            'companyName': company,
            'industry': industry,
            'employeeScale': scale,
            'phone': phone,
            'contactName': contact,
            'status': item_status,
            'memo': memo,
            'callHistory': call_history,
            'keepHistory': [],
            'isKept': False,
            'keptBy': '',
            'keptAt': '',
            'salesRep': '未確定',
            'email': '',
            'emailStatus': '未送信',
            'companyUrl': '',
            'prefecture': '',
            'listingStatus': '',
            'recruitmentPhone': '',
            'recruitmentAdSpend': '',
            'recruitmentJobType': '',
            'salesScale': '',
            'employeeGrowth': '',
            'nextCallDate': '',
            'listSource': sheet_name,
        }
        new_items.append(item)
        # 重複防止用に追加
        existing_names.add(company.lower())
        if phone: existing_phones.add(phone)

    print(f"インポート対象: {len(new_items)}件")
    print(f"スキップ（重複）: {skipped_dup}件 / 空行: {skipped_empty}件")

    if not new_items:
        print("インポートするデータがありません")
        return

    # サンプル表示
    s = new_items[0]
    print(f"\n--- サンプル1件目 ---")
    print(f"  企業名: {s['companyName']}")
    print(f"  業種: {s['industry']} / 規模: {s['employeeScale']}")
    print(f"  電話: {s['phone']} / 担当者: {s['contactName']}")
    print(f"  ステータス: {s['status']} / 架電履歴: {len(s['callHistory'])}件")
    print(f"  メモ: {s['memo'][:50] if s['memo'] else '(なし)'}")
    if s['callHistory']:
        c = s['callHistory'][-1]
        print(f"  最終架電: {c['calledAt']} by {c['calledBy']} → {c['result']}")

    if dry_run:
        print("\n[ドライラン] 実際のインポートはスキップしました")
        return new_items

    confirm = input(f"\n{len(new_items)}件をインポートしますか？ (y/n): ")
    if confirm.lower() != 'y':
        print("キャンセルしました")
        return

    # バッチ挿入
    BATCH = 50
    success = 0
    for i in range(0, len(new_items), BATCH):
        batch = new_items[i:i + BATCH]
        payload = json.dumps([{'id': item['id'], 'data': item} for item in batch]).encode()
        req = urllib.request.Request(
            SUPABASE_URL + '/rest/v1/teleapo_items',
            data=payload,
            headers={
                'apikey': SUPABASE_KEY,
                'Authorization': 'Bearer ' + SUPABASE_KEY,
                'Content-Type': 'application/json',
                'Prefer': 'return=minimal',
            },
            method='POST'
        )
        urllib.request.urlopen(req)
        success += len(batch)
        print(f"  {success}/{len(new_items)}件完了...")

    print(f"\n✅ {success}件のインポートが完了しました")

if __name__ == '__main__':
    # まずドライランで確認
    import_sheet('住宅・リフォーム', dry_run=True)
